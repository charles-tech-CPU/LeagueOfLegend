#!/usr/bin/env python3
"""
Convertit le fichier Excel de suivi LoL de Charles (feuille "Feuil1", le
calendrier a plat) en un script SQL (V2__seed_data.sql) que Flyway jouera
automatiquement au demarrage du backend.

Pourquoi seulement "Feuil1" et pas les feuilles par ligue (LCK, LPL, LEC...) ?
Ces dernieres sont une mise en forme VISUELLE (matrice tete-a-tete calculee
a la main + bracket de playoffs positionne librement dans les cellules), pas
une structure de donnees exploitable automatiquement. Feuil1 contient deja
tous les memes matchs sous une forme plate et coherente (BO, region, date,
heure, round, equipe1, equipe2, score1, score2), donc c'est la seule source
utilisee pour l'import. Les feuilles MSI / Esports World Cup ne sont pas
dans Feuil1 (evenements internationaux non repris dans le calendrier a
plat) : a saisir a la main via l'appli si besoin, ou a importer plus tard
avec un script dedie.

Usage :
    python3 import_excel.py /chemin/vers/LOL_2026.xlsx > ../backend/src/main/resources/db/migration/V2__seed_data.sql
"""
import sys
from datetime import datetime

import openpyxl

SHEET = "Feuil1"
SEASON = 2026


def sql_escape(value: str) -> str:
    return value.replace("'", "''")


def parse_time(raw):
    if not raw:
        return "NULL"
    # format observe dans le fichier : "11H00", "13H30" ...
    hh, mm = raw.strip().upper().split("H")
    return f"'{int(hh):02d}:{int(mm):02d}:00'"


def main():
    if len(sys.argv) != 2:
        print("Usage: import_excel.py <fichier.xlsx>", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    ws = wb[SHEET]

    teams = {}       # code -> region (region du premier match rencontre)
    regions = set()
    match_rows = []

    for row in range(4, ws.max_row + 1):
        bo = ws.cell(row=row, column=2).value
        region = ws.cell(row=row, column=3).value
        date = ws.cell(row=row, column=4).value
        horaire = ws.cell(row=row, column=5).value
        round_label = ws.cell(row=row, column=6).value
        t1 = ws.cell(row=row, column=7).value
        t2 = ws.cell(row=row, column=8).value
        s1 = ws.cell(row=row, column=9).value
        s2 = ws.cell(row=row, column=10).value

        if not region or not t1 or not t2 or not date:
            continue  # ligne vide / separatrice

        regions.add(region)
        teams.setdefault(t1, region)
        teams.setdefault(t2, region)

        match_rows.append({
            "bo": bo or "BO3",
            "region": region,
            "date": date if isinstance(date, datetime) else None,
            "time": horaire,
            "round": round_label or "?",
            "t1": t1,
            "t2": t2,
            "s1": s1,
            "s2": s2,
        })

    out = sys.stdout

    out.write("-- Genere automatiquement par import/import_excel.py depuis le fichier Excel de Charles.\n")
    out.write(f"-- {len(teams)} equipes, {len(regions)} competitions (saison {SEASON}), {len(match_rows)} matchs.\n\n")

    out.write("-- Equipes (code = nom pour l'instant ; a completer/renommer via l'appli si besoin).\n")
    for code, region in sorted(teams.items()):
        out.write(
            f"INSERT INTO team (code, name, region) VALUES "
            f"('{sql_escape(code)}', '{sql_escape(code)}', '{sql_escape(region)}');\n"
        )

    out.write("\n-- Competitions (une ligue = une competition pour la saison {}).\n".format(SEASON))
    for region in sorted(regions):
        out.write(
            f"INSERT INTO competition (code, name, type, region, season) VALUES "
            f"('{sql_escape(region)}', '{sql_escape(region)} {SEASON}', 'REGIONAL_LEAGUE', "
            f"'{sql_escape(region)}', {SEASON});\n"
        )

    out.write("\n-- Matchs (calendrier + resultats).\n")
    for m in match_rows:
        if m["date"] is None:
            continue
        date_sql = m["date"].strftime("%Y-%m-%d")
        time_sql = parse_time(m["time"])
        s1_sql = "NULL" if m["s1"] is None else int(m["s1"])
        s2_sql = "NULL" if m["s2"] is None else int(m["s2"])
        status = "COMPLETED" if m["s1"] is not None and m["s2"] is not None else "SCHEDULED"
        out.write(
            "INSERT INTO match (competition_id, round_label, date, time, best_of, "
            "team1_id, team2_id, score1, score2, status) VALUES ("
            f"(SELECT id FROM competition WHERE code = '{sql_escape(m['region'])}' AND season = {SEASON}), "
            f"'{sql_escape(str(m['round']))}', '{date_sql}', {time_sql}, '{sql_escape(m['bo'])}', "
            f"(SELECT id FROM team WHERE code = '{sql_escape(m['t1'])}'), "
            f"(SELECT id FROM team WHERE code = '{sql_escape(m['t2'])}'), "
            f"{s1_sql}, {s2_sql}, '{status}');\n"
        )

    print(
        f"OK : {len(teams)} equipes, {len(regions)} competitions, {len(match_rows)} matchs ecrits.",
        file=sys.stderr,
    )


if __name__ == "__main__":
    main()
